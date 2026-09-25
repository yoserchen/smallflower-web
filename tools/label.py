#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
小花老師的花曆 — 每月標註工具

用法（在 smallflower-web 資料夾裡執行）：
    python3 tools/label.py

它會去 ~/Documents/花曆更新/待標註/ 找照片，做成縮圖，
產生一個 ~/Documents/花曆標註/標註.html，用瀏覽器打開就能一張一張打花名。

標完按「匯出」，把下載到的「新增記錄.json」放到 ~/Documents/花曆更新/ 底下，
再跑 python3 tools/update.py 就會把新的月份、記錄數、新的花全部帶進網站。
"""
import os, sys, json, shutil, datetime, hashlib, re, glob, math, collections
import xml.etree.ElementTree as ET

try:
    from PIL import Image, ImageOps
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except ImportError:
        pass
except ImportError:
    print('少了 Pillow。請先執行：pip3 install --user pillow'); sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IN = os.path.expanduser('~/Documents/花曆更新/待標註')
OUT = os.path.expanduser('~/Documents/花曆標註')
SLUG = 'tpbg'
EXTS = ('.jpg', '.jpeg', '.png', '.heic', '.heif', '.tif', '.tiff', '.webp')
sys.path.insert(0, os.path.join(ROOT, 'tools'))

if not os.path.isdir(IN):
    os.makedirs(IN, exist_ok=True)
    print(f'已經建好 {IN}')
    print('把這個月拍的照片放進去，再執行一次。')
    sys.exit(0)

files = sorted(f for f in os.listdir(IN) if f.lower().endswith(EXTS) and not f.startswith('.'))
if not files:
    print(f'{IN} 裡沒有照片。'); sys.exit(0)
print(f'找到 {len(files)} 張照片')

# ---------- 既有的花名與區域 ----------
species = json.load(open(os.path.join(ROOT, 'src/data', SLUG + '.json'), encoding='utf-8'))
known = {}
for s in species:
    if s['z']:
        known.setdefault(s['n'], s['z'])
    else:
        known.setdefault(s['n'], '')
zones = sorted({s['z'] for s in species if s['z']})


# ---------- 用既有花點推算「這個座標屬於哪一區」 ----------
NS = '{http://www.opengis.net/kml/2.2}'
code_zone = {}
try:
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(ROOT, 'tools/data/主檔.xlsx'), data_only=True)
    it = wb['物種主檔'].iter_rows(values_only=True)
    hd = [str(x).strip() if x else '' for x in next(it)]
    ci = {k: i for i, k in enumerate(hd)}
    ZR = {'蜜源植物區': '植物與昆蟲區', '名人植物區': '植物名人園區'}
    for r in it:
        if r and r[0]:
            z = (r[ci['區域']] or '').strip()
            code_zone[str(r[ci['內部代號']]).strip()] = ZR.get(z, z)
except Exception:
    pass

ref = []          # [(lat, lon, 區域)]
for f in glob.glob(os.path.join(ROOT, 'tools/data/kml/*.kml')):
    try:
        root = ET.parse(f).getroot()
    except Exception:
        continue
    for pm in root.iter(NS + 'Placemark'):
        nm = (pm.findtext(NS + 'name') or '').strip()
        m = re.match(r'(\S+?-\d+)', nm)
        co = pm.find('.//' + NS + 'coordinates')
        if not m or co is None:
            continue
        z = code_zone.get(m.group(1))
        if not z:
            continue
        try:
            lo, la = co.text.strip().split(',')[:2]
            ref.append((float(la), float(lo), z))
        except Exception:
            pass


def metres(a, b):
    dy = (a[0] - b[0]) * 111320
    dx = (a[1] - b[1]) * 111320 * math.cos(math.radians(a[0]))
    return math.hypot(dx, dy)


def zone_of(lat, lon, k=9):
    """看最近的 k 個既有花點在哪一區，多數決。"""
    if not ref:
        return '', 0
    near = sorted(ref, key=lambda r: metres((lat, lon), (r[0], r[1])))[:k]
    vote = collections.Counter(r[2] for r in near)
    z, n = vote.most_common(1)[0]
    return z, round(metres((lat, lon), (near[0][0], near[0][1])))


def gps(path):
    try:
        g = Image.open(path).getexif().get_ifd(0x8825)
        if not g:
            return None
        def dms(v):
            d, m, s = [float(x) for x in v]
            return d + m / 60 + s / 3600
        lat, lon = dms(g[2]), dms(g[4])
        if g.get(1) == 'S':
            lat = -lat
        if g.get(3) == 'W':
            lon = -lon
        if not (24.9 < lat < 25.2 and 121.3 < lon < 121.7):
            return None          # 不在台北，多半是壞資料
        return round(lat, 7), round(lon, 7)
    except Exception:
        return None


def shot_time(path):
    try:
        ex = Image.open(path).getexif()
        for tag in (36867, 36868, 306):
            v = ex.get(tag)
            if v:
                return datetime.datetime.strptime(str(v)[:19], '%Y:%m:%d %H:%M:%S')
    except Exception:
        pass
    return datetime.datetime.fromtimestamp(os.path.getmtime(path))


# ---------- 縮圖 ----------
TH = os.path.join(OUT, 'thumbs')
FU = os.path.join(OUT, 'full')
for d in (TH, FU):
    shutil.rmtree(d, ignore_errors=True)
    os.makedirs(d)

photos = []
for i, f in enumerate(files, 1):
    src = os.path.join(IN, f)
    try:
        im = Image.open(src)
        im.draft('RGB', (1600, 1600))
        im = ImageOps.exif_transpose(im).convert('RGB')
    except Exception as e:
        print(f'  ! 讀不了 {f}：{e}')
        continue
    name = os.path.splitext(f)[0] + '.jpg'
    ImageOps.fit(im, (300, 300), Image.LANCZOS, centering=(.5, .45)) \
        .save(os.path.join(TH, name), 'JPEG', quality=62, optimize=True)
    big = im.copy(); big.thumbnail((1400, 1400), Image.LANCZOS)
    big.save(os.path.join(FU, name), 'JPEG', quality=80, optimize=True)
    t = shot_time(src)
    rec = {'f': name, 'src': f, 'd': t.strftime('%Y-%m-%d')}
    ll = gps(src)
    if ll:
        rec['ll'] = ll
        z, dd = zone_of(*ll)
        if z:
            rec['gz'] = z
            rec['gd'] = dd
    photos.append(rec)
    if i % 50 == 0:
        print(f'  處理 {i}/{len(files)}')

if not photos:
    print('沒有一張照片讀得進來。'); sys.exit(1)

batch = hashlib.md5((''.join(p['f'] for p in photos)).encode()).hexdigest()[:8]
blob = json.dumps({'batch': batch, 'photos': photos, 'species': known, 'zones': zones},
                  ensure_ascii=False, separators=(',', ':'))

from label_tpl import LABEL_HTML  # noqa: E402
with open(os.path.join(OUT, '標註.html'), 'w', encoding='utf-8') as fh:
    fh.write(LABEL_HTML.replace('__DATA__', blob))

# 記住這一批照片的原始檔名，之後歸檔用
json.dump({p['f']: p['src'] for p in photos},
          open(os.path.join(OUT, 'batch.json'), 'w', encoding='utf-8'), ensure_ascii=False)

ds = sorted({p['d'] for p in photos})
ng = sum(1 for p in photos if p.get('ll'))
print()
print(f'{len(photos)} 張，拍攝日期 {ds[0]} ～ {ds[-1]}')
print(f'其中 {ng} 張有 GPS' + ('，區域會自動帶入' if ng else '；沒有 GPS 的要自己選區域'))
print(f'用瀏覽器打開：{OUT}/標註.html')
print('標完按「匯出」，把「新增記錄.json」放到 ~/Documents/花曆更新/ 底下，再跑 tools/update.py')
